"""DHS Logistics: report emails (weekly report, monthly master file, daily contact list).

GitHub runs this every 30 minutes (see .github/workflows/report-emails.yml). Each time it reads the
settings an Admin keeps on the dashboard (Team page -> Report emails) and sends whatever is due:
- Weekly report: on the chosen weekday and time, the orders since the previous weekly report, as Excel.
- Monthly master file: on the chosen day of the month and time, every order of the previous month
  (DHS_Master_YYYY-MM.xlsx), the copy to keep in the master folder.
- Daily contact list: every day at the chosen time, the dealer, party, phone and location of every order
  from the 1st of the month until then (DHS_Contacts_YYYY-MM.xlsx). It grows through the month and starts
  afresh on the 1st.
Each email has its own on/off switch and its own list of recipients.
Test runs (started by hand with "Run workflow") send all three emails at once with [TEST] in the subject
and don't change what's due.

Secrets it needs (GitHub: Settings -> Secrets and variables -> Actions):
  FIREBASE_SERVICE_ACCOUNT  the Firebase service account key (the whole JSON file's text)
  GMAIL_USER                the Gmail address that sends the reports
  GMAIL_APP_PASSWORD        a Gmail app password for that address
"""
import io, json, os, smtplib, sys
from datetime import datetime, timedelta
from email.message import EmailMessage
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

IST = ZoneInfo("Asia/Kolkata")
COLUMNS = ["Order ID", "Date", "Order of the Day", "Order Date & Time", "Scheduled For", "Dealer", "Party", "Phone", "Location", "Products",
           "Total Bags", "Vehicle", "Driver", "Dispatch Depot", "Status", "Hold / Cancel Reason", "Cancelled By", "Cancelled At", "Dispatched Date & Time", "Time Taken (h:mm)",
           "Time Taken (minutes)", "Remarks", "Entered by", "Last Changed By", "Last Changed At"]


# ---------- helpers (same rules as the dashboard) ----------
def dt(ms):
    return datetime.fromtimestamp(ms / 1000, IST)

def day_key(ms):
    return dt(ms).strftime("%Y-%m-%d")

def iso_dt(ms):
    return "" if ms is None else dt(ms).strftime("%Y-%m-%d %H:%M")

def hm(ms):
    if ms is None:
        return ""
    m = max(0, round(ms / 60000))
    return f"{m // 60}:{m % 60:02d}"

def q_at(o):
    """When the order joins the queue: an advance order joins at the start of the day it is meant to go."""
    s = o.get("schedAt")
    return max(o["orderAt"], s) if isinstance(s, (int, float)) else o["orderAt"]

def status_of(o):
    if o.get("dispatchedAt") is not None:
        return "Dispatched"
    if o.get("mark") == "Cancelled":
        return "Cancelled"
    if o.get("hold") is True:
        return "On hold"
    s = o.get("schedAt")
    if isinstance(s, (int, float)) and day_key(s) > datetime.now(IST).strftime("%Y-%m-%d"):
        return "Scheduled"
    return "Pending"

def items_text(o):
    items = o.get("items") or ([{"product": o["brand"], "bags": o.get("bags", 0)}] if o.get("brand") else [])
    return ", ".join(f"{i.get('product', '')} × {i.get('bags', '')}" for i in items)

def with_tokens(orders):
    """Token = order's serial number within its day, by order time (whole days must be passed in)."""
    orders = sorted([o for o in orders if isinstance(o.get("orderAt"), (int, float))], key=lambda o: (q_at(o), o.get("createdAt") or 0, o["id"]))
    per = {}
    for o in orders:
        o["day"] = o.get("day") or day_key(o["orderAt"])
        per[o["day"]] = per.get(o["day"], 0) + 1
        o["token"] = per[o["day"]]
        o["status"] = status_of(o)
        o["taken"] = (o["dispatchedAt"] - q_at(o)) if o.get("dispatchedAt") is not None else None
    return orders

def agg(lst):
    a = dict(n=0, Pending=0, Hold=0, Dispatched=0, Cancelled=0, bags=0, pbags=0, hbags=0, cbags=0, tsum=0, tn=0)
    for o in lst:
        st = "Hold" if o["status"] == "On hold" else "Pending" if o["status"] == "Scheduled" else o["status"]     # on hold has its own head, apart from pending
        a["n"] += 1; a[st] += 1
        b = int(o.get("bags") or 0)
        if st == "Cancelled": a["cbags"] += b
        else: a["bags"] += b
        if st == "Pending": a["pbags"] += b
        if st == "Hold": a["hbags"] += b
        if o["taken"] is not None: a["tsum"] += o["taken"]; a["tn"] += 1
    a["avg"] = hm(a["tsum"] / a["tn"]) if a["tn"] else ""
    return a

def row(o, names):
    sch = o["day"] if isinstance(o.get("schedAt"), (int, float)) and day_key(o["schedAt"]) > day_key(o["orderAt"]) else ""
    return [o["id"], o["day"], o["token"], iso_dt(o["orderAt"]), sch, o.get("dealer", ""), o.get("party", ""), o.get("phone", ""),
            o.get("location", ""), items_text(o), int(o.get("bags") or 0), o.get("vehicle", ""), o.get("driver", ""), o.get("depot", "") or "", o["status"],
            (o.get("holdNote") or "") if o["status"] == "On hold" else (o.get("cancelNote") or "") if o["status"] == "Cancelled" else "",
            names.get(o.get("cancelledBy"), o.get("cancelledBy") or "") if o["status"] == "Cancelled" and o.get("cancelledBy") else "",
            iso_dt(o.get("cancelledAt")) if o["status"] == "Cancelled" else "",
            iso_dt(o.get("dispatchedAt")), hm(o["taken"]), round(o["taken"] / 60000) if o["taken"] is not None else "",
            o.get("remarks", ""), names.get(o.get("createdBy"), o.get("createdBy") or ""),
            names.get(o.get("updatedBy"), o.get("updatedBy") or "") if o.get("updatedBy") else "", iso_dt(o.get("updatedAt"))]


# ---------- Excel ----------
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial")

TIME_COLUMNS = {"Time Taken (h:mm)", "Time Taken (minutes)"}

def workbook(title, start_label, end_label, orders, names, with_time=True):
    """with_time=False leaves out the time-taken columns (per order, and the average on the Summary sheet)."""
    keep = [i for i, c in enumerate(COLUMNS) if with_time or c not in TIME_COLUMNS]
    wb = Workbook()
    ws = wb.active; ws.title = "Orders"
    ws.append([COLUMNS[i] for i in keep])
    for o in orders:
        r = row(o, names)
        ws.append([r[i] for i in keep])
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in ws.iter_rows(min_row=2):
        for c in r: c.font = BODY_FONT
    widths = [24, 11, 7, 17, 12, 22, 22, 13, 22, 30, 10, 13, 13, 15, 11, 24, 18, 17, 19, 11, 11, 22, 18, 18, 17]
    for n, i in enumerate(keep, 1): ws.column_dimensions[get_column_letter(n)].width = widths[i]
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions

    s = wb.create_sheet("Summary")
    s.append([f"DHS Logistics: {title}"]); s["A1"].font = Font(name="Arial", bold=True, size=13)
    s.append(["From", start_label, "To", end_label])
    s.append(["Prepared", datetime.now(IST).strftime("%Y-%m-%d %H:%M")])
    head = lambda first: [first, "Orders", "Pending", "On hold", "Dispatched", "Cancelled", "Bags", "Pending bags", "Cancelled bags"] + (["Avg time (h:mm)"] if with_time else [])
    line = lambda label, a: [label, a["n"], a["Pending"], a["Hold"], a["Dispatched"], a["Cancelled"], a["bags"], a["pbags"], a["cbags"]] + ([a["avg"]] if with_time else [])
    total = agg(orders)
    for heading, key, first in (("By dealer", lambda o: o.get("dealer") or "(no dealer)", "Dealer"), ("By day", lambda o: o["day"], "Date")):
        s.append([]); s.append([heading]); s.cell(s.max_row, 1).font = Font(name="Arial", bold=True)
        s.append(head(first))
        for c in s[s.max_row]: c.fill, c.font = HEAD_FILL, HEAD_FONT
        groups = {}
        for o in orders: groups.setdefault(key(o), []).append(o)
        for k in sorted(groups): s.append(line(k, agg(groups[k])))
        s.append(line("Total", total))
        for c in s[s.max_row]: c.font = Font(name="Arial", bold=True)
    s.column_dimensions["A"].width = 26
    for col in "BCDEFGHIJ": s.column_dimensions[col].width = 13
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue(), total


# ---------- data ----------
class FirestoreData:
    def __init__(self):
        import firebase_admin
        from firebase_admin import credentials, firestore
        info = json.loads(os.environ["FIREBASE_SERVICE_ACCOUNT"])
        firebase_admin.initialize_app(credentials.Certificate(info))
        self.db = firestore.client()

    def orders_between_days(self, d_from, d_to):
        q = self.db.collection("orders").where("day", ">=", d_from).where("day", "<=", d_to)
        return [{"id": d.id, **d.to_dict()} for d in q.stream()]

    def names(self):
        return {d.id: (d.to_dict() or {}).get("name") or d.id for d in self.db.collection("members").stream()}

    def report_settings(self):
        snap = self.db.collection("config").document("report").get()
        return snap.to_dict() if snap.exists else {}

    def save_report_settings(self, fields):
        self.db.collection("config").document("report").set(fields, merge=True)


# ---------- settings ----------
DAY_NAMES = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
DEFAULT_WEEKLY = {"on": True, "day": 5, "time": "21:00", "recipients": []}      # Saturday 9 pm
DEFAULT_MONTHLY = {"on": True, "day": 1, "time": "09:00", "recipients": []}     # 1st of the month, 9 am
DEFAULT_DAILY = {"on": True, "time": "21:00", "recipients": []}                 # every day, 9 pm

def settings(cfg):
    w = {**DEFAULT_WEEKLY, **(cfg.get("weekly") or {})}
    if not (cfg.get("weekly") or {}).get("recipients") and cfg.get("recipients"):
        w["recipients"] = cfg["recipients"]                     # list saved by an earlier version
    m = {**DEFAULT_MONTHLY, **(cfg.get("monthly") or {})}
    d = {**DEFAULT_DAILY, **(cfg.get("daily") or {})}
    for s in (w, m, d):
        s["recipients"] = [e.strip().lower() for e in (s.get("recipients") or []) if "@" in str(e)]
        hh, mm = (str(s.get("time") or "09:00") + ":00").split(":")[:2]
        s["hh"], s["mm"] = min(23, max(0, int(hh))), min(59, max(0, int(mm)))
    w["day"] = int(w.get("day", 5)) % 7
    m["day"] = min(28, max(1, int(m.get("day", 1))))
    return w, m, d

def last_weekly_moment(now, w):
    d = now.replace(hour=w["hh"], minute=w["mm"], second=0, microsecond=0)
    d -= timedelta(days=(d.weekday() - w["day"]) % 7)
    if d > now: d -= timedelta(days=7)
    return d

def this_month_moment(now, m):
    return now.replace(day=m["day"], hour=m["hh"], minute=m["mm"], second=0, microsecond=0)

def nice_time(s):
    return datetime(2000, 1, 1, s["hh"], s["mm"]).strftime("%I:%M %p").lstrip("0").lower()

def ms(d): return int(d.timestamp() * 1000)


# ---------- emails ----------
XLSX = ("application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet")

def make_email(subject, body, recipients, attachments):
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = f"DHS Logistics Reports <{os.environ.get('GMAIL_USER', 'reports@example.com')}>"
    msg["To"] = ", ".join(recipients)
    msg.set_content(body)
    for fname, content in attachments:
        msg.add_attachment(content, maintype=XLSX[0], subtype=XLSX[1], filename=fname)
    return msg

FOOTER = "Sent automatically by the DHS Logistics - Orders Dashboard (Alpha). An Admin can change who gets this email, and when, on the dashboard under Settings -> Team & emails."

def cancelled_lines(lst, names):
    """One line per cancelled order: what it was, who cancelled it and when, and why."""
    out = []
    for o in lst:
        if o.get("dispatchedAt") is not None or o.get("mark") != "Cancelled":
            continue
        who = f"cancelled by {names.get(o.get('cancelledBy'), o.get('cancelledBy'))}" if o.get("cancelledBy") else "cancelled (who cancelled wasn't recorded)"
        when = f" at {dt(o['cancelledAt']).strftime('%d %b, %I:%M %p')}" if isinstance(o.get("cancelledAt"), (int, float)) else ""
        why = f". Reason: {o['cancelNote']}" if o.get("cancelNote") else ""
        out.append(f"  {dt(o['orderAt']).strftime('%d %b')}: {o.get('party') or '?'} ({o.get('dealer') or 'no dealer'}, {int(o.get('bags') or 0)} bags), {who}{when}{why}")
    return "\n".join(out)

def weekly_email(data, w, start, end, test, names):
    s_ms, e_ms = ms(start), ms(end)
    full_days = with_tokens(data.orders_between_days(start.strftime("%Y-%m-%d"), end.strftime("%Y-%m-%d")))
    week = [o for o in full_days if s_ms <= o["orderAt"] < e_ms]
    label = f"{start.strftime('%d %b')} to {end.strftime('%d %b %Y')}"
    xl, t = workbook(f"Weekly report, {label}", iso_dt(s_ms), iso_dt(e_ms), week, names, with_time=False)
    groups = {}
    for o in week: groups.setdefault(o.get("dealer") or "(no dealer)", []).append(o)
    dealers = "\n".join(f"  {k}: {agg(v)['n']} orders, {agg(v)['bags']:,} bags" for k, v in sorted(groups.items(), key=lambda kv: -len(kv[1])))
    body = f"""DHS Logistics weekly report
{start.strftime('%a %d %b, %I:%M %p')} to {end.strftime('%a %d %b, %I:%M %p')}

Orders received: {t['n']}
Dispatched: {t['Dispatched']}
Still pending: {t['Pending']} ({t['pbags']:,} bags)
On hold: {t['Hold']} ({t['hbags']:,} bags)
Cancelled: {t['Cancelled']} ({t['cbags']:,} bags)
Bags (not counting cancelled): {t['bags']:,} (about {t['bags'] * 50 / 1000:,.1f} t)

By dealer:
{dealers or '  No orders in this period.'}

Cancelled orders and who cancelled them:
{cancelled_lines(week, names) or '  None.'}

The attached Excel has every order of the period (Orders sheet) and the totals (Summary sheet).

{FOOTER}
"""
    subject = ("[TEST] " if test else "") + f"DHS Logistics weekly report: {label}"
    fname = f"DHS_Week_{start.strftime('%Y-%m-%d')}_to_{end.strftime('%Y-%m-%d')}.xlsx"
    return make_email(subject, body, w["recipients"], [(fname, xl)]), t["n"]

def monthly_email(data, m, month, test, names):
    y, mo = map(int, month.split("-"))
    last_day = ((datetime(y + (mo == 12), mo % 12 + 1, 1)) - timedelta(days=1)).day
    orders = with_tokens(data.orders_between_days(f"{month}-01", f"{month}-{last_day:02d}"))
    if not orders and not test:
        return None, 0
    xl, t = workbook(f"Master file {month}", f"{month}-01", f"{month}-{last_day:02d}", orders, names)
    name = datetime(y, mo, 1).strftime("%B %Y")
    body = f"""DHS Logistics master file for {name}

Orders: {t['n']} (dispatched {t['Dispatched']}, still pending {t['Pending']}, on hold {t['Hold']}, cancelled {t['Cancelled']})
Cancelled bags: {t['cbags']:,}
Bags (not counting cancelled): {t['bags']:,} (about {t['bags'] * 50 / 1000:,.1f} t)
Average time to dispatch: {t['avg'] or '–'}

Attached: DHS_Master_{month}.xlsx with every order of {name}. Save it in your DHS Master folder
(replace the old copy if there is one), then click Data -> Refresh All in your master Excel.

{FOOTER}
"""
    subject = ("[TEST] " if test else "") + f"DHS Logistics master file: {name}"
    return make_email(subject, body, m["recipients"], [(f"DHS_Master_{month}.xlsx", xl)]), t["n"]


CONTACT_COLUMNS = ["Dealer", "Party", "Phone", "Location"]

def contacts_workbook(orders):
    wb = Workbook()
    ws = wb.active; ws.title = "Contacts"
    ws.append(CONTACT_COLUMNS)
    for o in orders:
        ws.append([o.get("dealer", "") or "", o.get("party", "") or "", o.get("phone", "") or "", o.get("location", "") or ""])
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(horizontal="center", vertical="center")
    for r in ws.iter_rows(min_row=2):
        for c in r: c.font = BODY_FONT
    for col, w in zip("ABCD", (26, 28, 15, 34)): ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()

def daily_email(data, d, now, test, names=None):
    """Every order from the 1st of this month up to now: dealer, party, phone, location (oldest first)."""
    first = now.replace(day=1)
    lst = data.orders_between_days(first.strftime("%Y-%m-%d"), now.strftime("%Y-%m-%d"))
    lst = sorted([o for o in lst if isinstance(o.get("orderAt"), (int, float)) and o["orderAt"] <= ms(now)],
                 key=lambda o: (o["orderAt"], o.get("createdAt") or 0, o["id"]))
    month = now.strftime("%Y-%m")
    label = f"1 to {now.day} {now.strftime('%b %Y')}" if now.day > 1 else f"1 {now.strftime('%b %Y')}"
    def counts(l):
        c = {"Dispatched": 0, "Pending": 0, "On hold": 0, "Scheduled": 0, "Cancelled": 0}
        for o in l: c[status_of(o)] += 1
        return f"{len(l)} (dispatched {c['Dispatched']}, pending {c['Pending']}, on hold {c['On hold']}, cancelled {c['Cancelled']}" + (f", for later days {c['Scheduled']}" if c["Scheduled"] else "") + ")"
    today = now.strftime("%Y-%m-%d")
    todays = [o for o in lst if (o.get("day") or day_key(o["orderAt"])) == today]
    body = f"""DHS Logistics daily contact list, {label}

Orders so far this month: {counts(lst)}
Orders today, {now.strftime('%d %b')}: {counts(todays)}
Cancelled today:
{cancelled_lines(todays, names or {}) or '  None.'}

Attached: DHS_Contacts_{month}.xlsx with the dealer, party, phone number and location of every order
from the 1st of the month until {now.strftime('%I:%M %p').lstrip('0').lower()} today. Each day's file has
everything since the 1st, so you can save it over yesterday's copy. It starts afresh on the 1st of next month.

{FOOTER}
"""
    subject = ("[TEST] " if test else "") + f"DHS Logistics daily contact list: {label} ({len(lst)} orders)"
    return make_email(subject, body, d["recipients"], [(f"DHS_Contacts_{month}.xlsx", contacts_workbook(lst))]), len(lst)


# ---------- the run ----------
def run(data, now=None, test=False, send=None):
    """Send whatever is due. Returns a list of what was sent (for the log and tests)."""
    now = now or datetime.now(IST)
    send = send or send_gmail
    cfg = data.report_settings()
    w, m, d = settings(cfg)
    names = None
    done, fields = [], {"lastCheckedAt": ms(now)}

    # Weekly report
    moment = now if test else last_weekly_moment(now, w)
    prev_end = cfg.get("lastWindowEnd")
    if not test and not prev_end:
        fields["lastWindowEnd"] = ms(moment)            # first run: start counting from the most recent scheduled time
        print("Weekly report: first check, counting starts from", moment.strftime("%a %d %b %I:%M %p"))
    elif w["on"] and w["recipients"] and (test or prev_end < ms(moment)):
        start = datetime.fromtimestamp(prev_end / 1000, IST) if (prev_end and not test and ms(moment) - prev_end <= 21 * 864e5) else moment - timedelta(days=7)
        names = names or data.names()
        msg, n = weekly_email(data, w, start, moment, test, names)
        send(msg); done.append(("weekly", n, list(w["recipients"])))
        fields["lastWeeklySentAt"] = ms(datetime.now(IST))
        fields["lastWeeklyResult"] = ("Test sent" if test else "Sent") + f" to {len(w['recipients'])}, {n} orders"
        if not test: fields["lastWindowEnd"] = ms(moment)
    elif not test and prev_end < ms(moment):
        fields["lastWindowEnd"] = ms(moment)            # switched off or no recipients: skip this period quietly
        fields["lastWeeklyResult"] = "Skipped: " + ("switched off" if not w["on"] else "no recipients on the list")

    # Monthly master file
    due = this_month_moment(now, m)
    month = (now.replace(day=1) - timedelta(days=1)).strftime("%Y-%m")       # the month before this one
    last_master = cfg.get("lastMasterMonth")
    if not test and not last_master:
        fields["lastMasterMonth"] = month              # first run: the first file will be for the next full month
        print("Monthly master: first check, the first file will cover the month after", month)
    elif m["on"] and m["recipients"] and (test or (now >= due and last_master < month)):
        names = names or data.names()
        msg, n = monthly_email(data, m, month, test, names)
        if msg:
            send(msg); done.append(("monthly", n, list(m["recipients"])))
            fields["lastMonthlySentAt"] = ms(datetime.now(IST))
            fields["lastMonthlyResult"] = ("Test sent" if test else "Sent") + f" to {len(m['recipients'])}, {month}, {n} orders"
        else:
            fields["lastMonthlyResult"] = f"Skipped: no orders in {month}"
        if not test: fields["lastMasterMonth"] = month
    elif not test and now >= due and last_master < month:
        fields["lastMasterMonth"] = month
        fields["lastMonthlyResult"] = "Skipped: " + ("switched off" if not m["on"] else "no recipients on the list")

    # Daily contact list (once a day, after the chosen time; month to date)
    today = now.strftime("%Y-%m-%d")
    due_daily = now.replace(hour=d["hh"], minute=d["mm"], second=0, microsecond=0)
    if d["on"] and d["recipients"] and (test or (now >= due_daily and cfg.get("lastDailyDate") != today)):
        names = names or data.names()
        msg, n = daily_email(data, d, now, test, names)
        send(msg); done.append(("daily", n, list(d["recipients"])))
        fields["lastDailySentAt"] = ms(datetime.now(IST))
        fields["lastDailyResult"] = ("Test sent" if test else "Sent") + f" to {len(d['recipients'])}, {n} orders this month"
        if not test: fields["lastDailyDate"] = today
    elif not test and now >= due_daily and cfg.get("lastDailyDate") != today:
        fields["lastDailyDate"] = today
        fields["lastDailyResult"] = "Skipped: " + ("switched off" if not d["on"] else "no recipients on the list")

    data.save_report_settings(fields)
    print("Sent:" if done else "Nothing due.", done)
    return done

def send_gmail(msg):
    user, pw = os.environ["GMAIL_USER"], os.environ["GMAIL_APP_PASSWORD"].replace(" ", "")
    with smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=60) as s:
        s.login(user, pw)
        s.send_message(msg)

if __name__ == "__main__":
    test = os.environ.get("TEST_RUN") == "1"
    data = FirestoreData()
    try:
        done = run(data, test=test)
        if done and os.environ.get("GITHUB_OUTPUT"):
            with open(os.environ["GITHUB_OUTPUT"], "a") as f: f.write("sent=yes\n")
    except Exception as e:
        print("Report email failed:", repr(e))
        try: data.save_report_settings({"lastError": "Failed: " + str(e)[:200], "lastCheckedAt": ms(datetime.now(IST))})
        except Exception: pass
        raise
